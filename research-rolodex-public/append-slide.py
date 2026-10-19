#!/usr/bin/env python3
"""
append-slide.py — Add a new paper slide (and section divider if needed) to an
existing Research Rolodex .pptx WITHOUT touching any existing slides.

Usage:
    python3 append-slide.py <config.json>

First run? Run setup.py first to create your deck and directory structure.

PSI branding assets are in the assets/ folder alongside this script.
To use a different logo, replace assets/psi-mark.png with your own PNG
(~0.45" × 0.34" aspect ratio works best), or set PSI_MARK_PNG = None
to show text-only footer.

The config.json for each paper specifies:
  {
    "deck_path":  "/path/to/Research Rolodex.pptx",   # required
    "theme":      "Machine Learning",                  # section name
    "layout":     "nofig",                             # nofig | wide | right | figs_only
    "paper": {
      "title":    "Full paper title here",
      "date":     "Jan 2024",
      "authors":  "Smith J, Jones A, et al.",
      "doi":      "10.1234/example",
      "rating":   8,
      "citation": "Nature Medicine, 2024",
      "body": [
        {"label": "Problem",    "text": "One sentence."},
        {"label": "Methods",    "text": "One sentence."},
        {"label": "Results",    "text": "Two sentences with key numbers."},
        {"label": "Conclusion", "text": "One sentence."}
      ],
      "notes": "Upload Date: Jan 2024\\n\\nMy Notes:\\n• \\n\\nHighlights:\\n• Key finding here."
    }
  }
"""
import json
import os
import sys
from copy import deepcopy
from pathlib import Path
import lxml.etree as etree

try:
    from PIL import Image as PILImage
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    from send2trash import send2trash
    HAS_SEND2TRASH = True
except ImportError:
    HAS_SEND2TRASH = False

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── CONSTANTS ──────────────────────────────────────────────────────────────────

# PSI logo mark — path relative to this script's directory.
# Set to None to show text-only footer.
_SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
PSI_MARK_PNG = os.path.join(_SCRIPT_DIR, "assets", "psi-mark.png")

# Theme display order — edit or extend freely; controls divider insertion order.
# Any theme not listed here is appended after these in first-appearance order.
THEME_ORDER: list[str] = []   # leave empty to use order of first appearance

# Six-color section palette, wraps for 7+ sections
# (section 7 reuses color 1, section 8 reuses color 2, etc.)
_PALETTE = [
    "0085CA",   # Blue
    "006747",   # Green
    "5D2584",   # Purple
    "C8102E",   # Red
    "007B8F",   # Teal
    "002D62",   # Navy
]

# Base colors
COLOR_ACCENT   = "221F72"   # PSI purple — default accent
COLOR_INK      = "141414"   # Body text
COLOR_GRAPHITE = "58595B"   # Secondary / labels
COLOR_PAPER    = "F7F7F4"   # Slide background (warm off-white)
COLOR_WHITE    = "FFFFFF"

# Fonts — Arial Narrow ships with every copy of Microsoft Office (Track B).
FONT_TITLE = "Arial Narrow"
FONT_BODY  = "Arial Narrow"
FONT_MONO  = "Consolas"
FONT       = FONT_BODY

# PSI phase-recognition bar — 6 segments, proportions LOCKED.
# (lavender / purple / magenta / cyan-deep / purple-70% / lavender)
# Fractions are of a 1920-unit-wide slide.
_PHASE_SEGS = [
    (0/1920,        220.38/1920, "908FB8"),
    (222.38/1920,   661.15/1920, "221F72"),
    (885.54/1920,   146.92/1920, "D80B8C"),
    (1034.46/1920,  183.65/1920, "009DD6"),
    (1220.12/1920,  514.23/1920, "221F72", 178),   # purple at 70% opacity
    (1736.35/1920,  183.65/1920, "908FB8"),
]


# ── COLOR HELPERS ─────────────────────────────────────────────────────────────

def section_color(theme: str, ordered_themes: list[str]) -> str:
    """Return a hex color for `theme`, wrapping the palette for 7+ sections.

    `ordered_themes` is built in main() as: THEME_ORDER + any existing deck
    sections not already listed + the incoming theme if new.  The position
    of `theme` in that list, taken mod 6, selects the color palette slot.
    """
    if theme in ordered_themes:
        idx = ordered_themes.index(theme)
        return _PALETTE[idx % len(_PALETTE)]
    return _PALETTE[0]


def hex_to_rgb(h: str) -> RGBColor:
    """Convert '7B9EB8' → RGBColor."""
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ── SLIDE ELEMENT HELPERS ──────────────────────────────────────────────────────

def add_phase_bar(slide, y_in: float = 0.0, h_in: float = 0.10,
                  slide_w_in: float = 10.0):
    """Add the PSI 6-segment phase-recognition bar across the slide top.
    Segment proportions are locked — do not reorder or recolor."""
    from pptx.oxml.ns import qn
    bar_top = Inches(y_in)
    bar_h   = Inches(h_in)
    slide_w = Inches(slide_w_in)
    for seg in _PHASE_SEGS:
        x_frac, w_frac, color_hex = seg[0], seg[1], seg[2]
        alpha = seg[3] if len(seg) == 4 else None
        seg_l = int(slide_w * x_frac)
        seg_w = int(slide_w * w_frac)
        if seg_w < 1:
            continue
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, seg_l, bar_top, seg_w, bar_h)
        rect.line.fill.background()
        rect.fill.solid()
        rect.fill.fore_color.rgb = hex_to_rgb(color_hex)
        if alpha is not None:
            solid = rect._element.find('.//' + qn('a:solidFill'))
            if solid is not None:
                clr = solid.find(qn('a:srgbClr'))
                if clr is not None:
                    al = etree.SubElement(clr, qn('a:alpha'))
                    al.set('val', str(int(alpha / 255 * 100000)))


def add_textbox(slide, text: str, x, y, w, h,
                font_size=11, bold=False, color="333333",
                align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
                font_face=FONT):
    """Add a simple single-run text box."""
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = hex_to_rgb(color)
    run.font.name = font_face
    txBox.text_frame._txBody.attrib[
        '{http://schemas.openxmlformats.org/drawingml/2006/main}anchor'
    ] = {MSO_ANCHOR.TOP: 't', MSO_ANCHOR.MIDDLE: 'ctr', MSO_ANCHOR.BOTTOM: 'b'}.get(valign, 't')
    return txBox


def add_rating_badge(slide, rating, color_hex: str = None):
    """Rating badge at top-left: section-color fill, white bold number."""
    from pptx.oxml.ns import qn
    badge = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.08), Inches(0.27), Inches(0.58), Inches(0.65))
    badge.fill.solid()
    badge.fill.fore_color.rgb = hex_to_rgb(color_hex or COLOR_ACCENT)
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = False
    bodyPr = tf._txBody.find(qn('a:bodyPr'))
    if bodyPr is not None:
        bodyPr.set('anchor', 'ctr')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(rating)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = hex_to_rgb(COLOR_WHITE)
    run.font.name = FONT_TITLE
    return badge


def add_rich_body(slide, sections, x, y, w, h, label_color: str = None):
    """Body content: bold label (section color) + regular body text."""
    label_color = label_color or COLOR_ACCENT
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, sec in enumerate(sections):
        if i > 0:
            spacer = tf.add_paragraph()
            spacer.space_before = Pt(2)
            spacer.space_after = Pt(2)
            sr = spacer.add_run()
            sr.text = ""
            sr.font.size = Pt(5)
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        label_run = p.add_run()
        label_run.text = sec["label"] + ": "
        label_run.font.size = Pt(11)
        label_run.font.bold = True
        label_run.font.color.rgb = hex_to_rgb(label_color)
        label_run.font.name = FONT_BODY
        text_run = p.add_run()
        text_run.text = sec["text"]
        text_run.font.size = Pt(11)
        text_run.font.bold = False
        text_run.font.color.rgb = hex_to_rgb(COLOR_INK)
        text_run.font.name = FONT_BODY
    return txBox


def add_citation_line(slide, opts, x=0.78, y=0.90, w=9.14, h=0.22):
    """Journal/citation line in Consolas graphite."""
    citation = opts.get("citation") or opts.get("date", "")
    if citation:
        add_textbox(slide, citation, x=x, y=y, w=w, h=h,
                    font_size=9, color=COLOR_GRAPHITE, font_face=FONT_MONO)


def add_psi_footer(slide):
    """PSI mark + 'Program for Surgical Intelligence' at bottom-left."""
    if PSI_MARK_PNG and os.path.exists(PSI_MARK_PNG):
        slide.shapes.add_picture(PSI_MARK_PNG,
                                 Inches(0.12), Inches(5.26),
                                 Inches(0.45), Inches(0.34))
    add_textbox(slide, "Program for Surgical Intelligence",
                x=0.62, y=5.32, w=3.2, h=0.22,
                font_size=7, color=COLOR_GRAPHITE, font_face=FONT_MONO)


def add_notes(slide, notes_text: str, prs=None):
    """Add speaker notes to a slide."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    if tf is not None:
        try:
            tf.text = notes_text
            return
        except Exception:
            pass
    for shape in notes_slide.shapes:
        if 'Notes Placeholder' in shape.name and shape.has_text_frame:
            shape.text_frame.text = notes_text
            return
    if prs is not None:
        from pptx.oxml.ns import qn
        from pptx.oxml import parse_xml
        template_sp = None
        for other_slide in prs.slides:
            if other_slide == slide:
                continue
            try:
                other_ns = other_slide.notes_slide
                for sp in other_ns._element.find(qn('p:cSld')).find(qn('p:spTree')).findall(qn('p:sp')):
                    nvSpPr = sp.find(qn('p:nvSpPr'))
                    if nvSpPr is not None:
                        cNvPr = nvSpPr.find(qn('p:cNvPr'))
                        if cNvPr is not None and 'Notes' in cNvPr.get('name', ''):
                            template_sp = sp
                            break
                if template_sp is not None:
                    break
            except Exception:
                continue
        if template_sp is not None:
            new_sp = deepcopy(template_sp)
            a_ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            txBody = new_sp.find('{http://schemas.openxmlformats.org/presentationml/2006/main}txBody') \
                     or new_sp.find(f'{{{a_ns}}}txBody')
            if txBody is not None:
                for p_elem in txBody.findall(f'{{{a_ns}}}p'):
                    txBody.remove(p_elem)
                escaped = notes_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                p_xml = f'<a:p xmlns:a="{a_ns}"><a:r><a:rPr lang="en-US" dirty="0"/><a:t>{escaped}</a:t></a:r></a:p>'
                txBody.append(parse_xml(p_xml))
            spTree = notes_slide._element.find(qn('p:cSld')).find(qn('p:spTree'))
            spTree.append(new_sp)
            return
    print("  WARNING: Could not add notes")


# ── SLIDE BUILDERS ─────────────────────────────────────────────────────────────

def create_title_slide(prs, title: str = "Research Rolodex",
                       subtitle: str = ""):
    """Create the cover/title slide. Called once by setup.py."""
    blank_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = hex_to_rgb(COLOR_PAPER)
    add_phase_bar(slide, h_in=0.15)
    add_textbox(slide, title,
                x=0.5, y=1.8, w=9.0, h=1.2,
                font_size=48, bold=True, color=COLOR_ACCENT,
                align=PP_ALIGN.CENTER, font_face=FONT_TITLE)
    if subtitle:
        add_textbox(slide, subtitle,
                    x=0.5, y=3.1, w=9.0, h=0.5,
                    font_size=18, color=COLOR_GRAPHITE,
                    align=PP_ALIGN.CENTER, font_face=FONT_BODY)
    return slide


def create_divider(prs, theme_name: str, color_hex: str = None):
    """Create a section divider slide."""
    color_hex = color_hex or COLOR_ACCENT
    blank_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = hex_to_rgb(COLOR_PAPER)
    add_phase_bar(slide)
    add_textbox(slide, theme_name,
                x=0, y=0, w=10, h=5.625,
                font_size=36, bold=True, color=color_hex,
                align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
                font_face=FONT_TITLE)
    return slide


def _paper_header(slide, opts: dict) -> str:
    """Shared header for all paper layouts: background, phase bar, badge, title.
    Returns the resolved section color hex string.
    """
    sc = opts.get("color") or COLOR_ACCENT

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = hex_to_rgb(COLOR_PAPER)

    add_phase_bar(slide)

    if opts.get("rating") is not None:
        add_rating_badge(slide, opts["rating"], sc)

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.78), Inches(0.27), Inches(9.14), Inches(0.50))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = opts["title"]
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(sc)
    r.font.name = FONT_TITLE

    if opts.get("authors"):
        add_textbox(slide, opts["authors"],
                    x=0.78, y=0.78, w=9.14, h=0.25,
                    font_size=10, color=COLOR_INK, font_face=FONT_BODY)

    add_citation_line(slide, opts, y=1.05)
    add_psi_footer(slide)

    return sc


def create_paper_slide_nofig(prs, opts: dict):
    """Text-only paper slide (no figure)."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    sc = _paper_header(slide, opts)
    add_rich_body(slide, opts["body"], x=0.50, y=1.20, w=9.0, h=3.75, label_color=sc)
    if opts.get("doi"):
        add_textbox(slide, "DOI: " + opts["doi"],
                    x=0.5, y=5.35, w=9.0, h=0.2,
                    font_size=8, color=COLOR_GRAPHITE, font_face=FONT_MONO)
    if opts.get("notes"):
        add_notes(slide, opts["notes"], prs)
    return slide


def create_paper_slide_wide(prs, opts: dict):
    """Paper slide with figure below body (full width)."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    sc = _paper_header(slide, opts)
    body_y = opts.get("bodyY", 1.20)
    body_h = opts.get("bodyHeight", 1.85)
    add_rich_body(slide, opts["body"], x=0.50, y=body_y, w=9.0, h=body_h, label_color=sc)
    if opts.get("figurePath"):
        fig_w = opts.get("figW", 6.2)
        aspect = opts["figAspectH"] / opts["figAspectW"]
        fig_h = fig_w * aspect
        fig_x = (10 - fig_w) / 2
        fig_y = opts.get("figY", 3.10)
        slide.shapes.add_picture(opts["figurePath"],
                                 Inches(fig_x), Inches(fig_y),
                                 Inches(fig_w), Inches(fig_h))
    if opts.get("doi"):
        add_textbox(slide, "DOI: " + opts["doi"],
                    x=0.5, y=5.35, w=9.0, h=0.2,
                    font_size=8, color=COLOR_GRAPHITE, font_face=FONT_MONO)
    if opts.get("notes"):
        add_notes(slide, opts["notes"], prs)
    return slide


def create_paper_slide_right(prs, opts: dict):
    """Paper slide with figure on the right."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    sc = _paper_header(slide, opts)
    body_y = opts.get("bodyY", 1.20)
    body_w = opts.get("bodyW", 5.0)
    body_h = opts.get("bodyHeight", 3.5)
    add_rich_body(slide, opts["body"], x=0.50, y=body_y, w=body_w, h=body_h, label_color=sc)
    if opts.get("figurePath"):
        fig_w = opts.get("figW", 3.8)
        aspect = opts["figAspectH"] / opts["figAspectW"]
        fig_h = fig_w * aspect
        fig_x = opts.get("figX", 5.9)
        fig_y = opts.get("figY", 0.80)
        slide.shapes.add_picture(opts["figurePath"],
                                 Inches(fig_x), Inches(fig_y),
                                 Inches(fig_w), Inches(fig_h))
    if opts.get("doi"):
        add_textbox(slide, "DOI: " + opts["doi"],
                    x=0.5, y=5.35, w=9.0, h=0.2,
                    font_size=8, color=COLOR_GRAPHITE, font_face=FONT_MONO)
    if opts.get("notes"):
        add_notes(slide, opts["notes"], prs)
    return slide


def create_paper_slide_figs_only(prs, opts: dict):
    """Figures-only companion slide (for papers with 2+ screenshots)."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _paper_header(slide, opts)
    fig_area = (0.15, 0.82, 9.7, 4.35)
    figure_paths = opts.get("figurePaths", [])
    n = len(figure_paths)
    fig_area_x, fig_area_y, fig_area_w, fig_area_h = fig_area
    if n == 1:
        cells = [(fig_area_x, fig_area_y, fig_area_w, fig_area_h)]
    elif n == 2:
        cw = (fig_area_w - 0.2) / 2
        cells = [(fig_area_x, fig_area_y, cw, fig_area_h),
                 (fig_area_x + cw + 0.2, fig_area_y, cw, fig_area_h)]
    elif n == 3:
        cw = (fig_area_w - 0.4) / 3
        cells = [(fig_area_x, fig_area_y, cw, fig_area_h),
                 (fig_area_x + cw + 0.2, fig_area_y, cw, fig_area_h),
                 (fig_area_x + 2*(cw + 0.2), fig_area_y, cw, fig_area_h)]
    else:
        cw = (fig_area_w - 0.2) / 2
        ch = (fig_area_h - 0.15) / 2
        cells = [(fig_area_x, fig_area_y, cw, ch),
                 (fig_area_x + cw + 0.2, fig_area_y, cw, ch),
                 (fig_area_x, fig_area_y + ch + 0.15, cw, ch),
                 (fig_area_x + cw + 0.2, fig_area_y + ch + 0.15, cw, ch)]
    for i, path in enumerate(figure_paths[:len(cells)]):
        if not os.path.exists(path):
            print(f"  WARNING: figure not found: {path}")
            continue
        cell_x, cell_y, cell_w, cell_h = cells[i]
        img_aspect = 1.0
        if HAS_PIL:
            try:
                with PILImage.open(path) as img:
                    iw, ih = img.size
                    img_aspect = ih / iw if iw > 0 else 1.0
            except Exception:
                pass
        if img_aspect * cell_w <= cell_h:
            fw, fh = cell_w, cell_w * img_aspect
        else:
            fh = cell_h
            fw = fh / img_aspect if img_aspect > 0 else cell_w
        fx = cell_x + (cell_w - fw) / 2
        fy = cell_y + (cell_h - fh) / 2
        slide.shapes.add_picture(path, Inches(fx), Inches(fy), Inches(fw), Inches(fh))
    if opts.get("doi"):
        add_textbox(slide, "DOI: " + opts["doi"],
                    x=0.5, y=5.35, w=9.0, h=0.2,
                    font_size=8, color=COLOR_GRAPHITE, font_face=FONT_MONO)
    if opts.get("notes"):
        add_notes(slide, opts["notes"], prs)
    return slide


# ── DECK ANALYSIS ──────────────────────────────────────────────────────────────

def find_theme_sections(prs) -> dict:
    """Scan existing slides and return section positions.

    Returns { theme_name: { 'divider_idx': int, 'last_paper_idx': int } }

    Divider detection: a slide with no DOI text and no substantial body text
    (>60 chars, positioned between 1" and 5" from top) is treated as a divider.
    The shortest non-empty text block on such a slide becomes the theme name.
    """
    sections = {}
    current_theme = None

    for idx, slide in enumerate(prs.slides):
        texts = [s.text_frame.text.strip()
                 for s in slide.shapes
                 if s.has_text_frame and s.text_frame.text.strip()]

        has_doi  = any('DOI' in t for t in texts)
        has_body = any(
            1.0 < s.top / 914400 < 5.0
            and s.has_text_frame
            and len(s.text_frame.text.strip()) > 60
            for s in slide.shapes
        )

        if not has_doi and not has_body:
            # Pick the shortest non-trivial text as the section name
            candidates = [t for t in texts if 2 < len(t) < 100 and 'DOI' not in t]
            if candidates:
                current_theme = min(candidates, key=len)
                sections[current_theme] = {
                    'divider_idx': idx,
                    'last_paper_idx': idx
                }
        elif current_theme:
            sections[current_theme]['last_paper_idx'] = idx

    return sections


def move_slide(prs, old_idx: int, new_idx: int):
    """Move slide from old_idx to new_idx in the presentation."""
    ns = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    sldIdLst = prs.part._element.find(f'{ns}sldIdLst')
    el = list(sldIdLst)[old_idx]
    sldIdLst.remove(el)
    sldIds = list(sldIdLst)
    if new_idx >= len(sldIds):
        sldIdLst.append(el)
    else:
        sldIds[new_idx].addprevious(el)


# ── MAIN ───────────────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 append-slide.py <config.json>")
        print("       python3 setup.py        ← run this first on a fresh install")
        sys.exit(1)

    with open(sys.argv[1]) as f:
        config = json.load(f)

    deck_path = config.get("deck_path")
    if not deck_path:
        print("ERROR: config.json must include 'deck_path'.")
        print("Run setup.py first to create your deck and generate a base config.")
        sys.exit(1)

    # Open or create deck
    if os.path.exists(deck_path):
        prs = Presentation(deck_path)
        print(f"Opened: {deck_path} ({len(prs.slides)} slides)")
    else:
        print(f"Deck not found at {deck_path}.")
        print("Run setup.py to create your deck first.")
        sys.exit(1)

    # Resolve theme and color
    theme   = config["theme"]
    sections = find_theme_sections(prs)
    print(f"Existing sections: {list(sections.keys())}")

    ordered_themes = list(THEME_ORDER)
    for t in sections:
        if t not in ordered_themes:
            ordered_themes.append(t)
    if theme not in ordered_themes:
        ordered_themes.append(theme)

    color        = section_color(theme, ordered_themes)
    footer_text  = config.get("footer_text", "")
    needs_divider = theme not in sections

    if needs_divider:
        print(f"Creating new section divider: {theme}")
        create_divider(prs, theme, color)
        new_divider_idx = len(prs.slides) - 1

        # Insert divider in correct THEME_ORDER position
        theme_idx = THEME_ORDER.index(theme) if theme in THEME_ORDER else len(THEME_ORDER)
        insert_before = None
        for later_theme in THEME_ORDER[theme_idx + 1:]:
            if later_theme in sections:
                insert_before = sections[later_theme]['divider_idx']
                break
        if insert_before is not None:
            move_slide(prs, new_divider_idx, insert_before)
            sections[theme] = {'divider_idx': insert_before, 'last_paper_idx': insert_before}
        else:
            sections[theme] = {'divider_idx': new_divider_idx, 'last_paper_idx': new_divider_idx}

    # Build paper opts
    layout      = config.get("layout", "nofig")
    paper_opts  = config["paper"]
    paper_opts["color"]       = color
    paper_opts["theme"]       = theme
    paper_opts["footer_text"] = footer_text

    layout_map = {
        "nofig":          create_paper_slide_nofig,
        "wide":           create_paper_slide_wide,
        "right":          create_paper_slide_right,
        "figs_only":      create_paper_slide_figs_only,
    }
    builder = layout_map.get(layout)
    if builder is None:
        raise ValueError(f"Unknown layout '{layout}'. Valid: {list(layout_map)}")
    builder(prs, paper_opts)
    new_slide_idx = len(prs.slides) - 1

    # Insert slide in the correct position within the theme section
    target_idx = sections[theme]['last_paper_idx'] + 1
    theme_idx = THEME_ORDER.index(theme) if theme in THEME_ORDER else len(THEME_ORDER)
    for later_theme in THEME_ORDER[theme_idx + 1:]:
        if later_theme in sections:
            target_idx = min(target_idx, sections[later_theme]['divider_idx'])
            break
    if new_slide_idx != target_idx:
        move_slide(prs, new_slide_idx, target_idx)
        print(f"  Inserted at position {target_idx}")

    # Companion figures-only slide
    embedded_figures = [paper_opts[k] for k in ("figurePath", "figurePath1", "figurePath2")
                        if paper_opts.get(k)]
    if "second_slide" in config:
        second_opts = dict(config["second_slide"])
        second_opts.update(color=color, theme=theme, footer_text=footer_text)
        for key in ("title", "date", "doi", "notes", "rating", "authors"):
            second_opts.setdefault(key, paper_opts.get(key))
        create_paper_slide_figs_only(prs, second_opts)
        second_idx = len(prs.slides) - 1
        move_slide(prs, second_idx, target_idx + 1)
        for path in second_opts.get("figurePaths", []):
            if path not in embedded_figures:
                embedded_figures.append(path)

    # Save
    output_path = config.get("output_path", deck_path)
    prs.save(output_path)
    print(f"Saved: {output_path} ({len(prs.slides)} slides)")

    # Move embedded figures to Trash
    if embedded_figures:
        if HAS_SEND2TRASH:
            for path in embedded_figures:
                if os.path.exists(path):
                    try:
                        send2trash(path)
                        print(f"  Trashed: {os.path.basename(path)}")
                    except Exception as e:
                        print(f"  Could not trash {path}: {e}")
        else:
            print("  NOTE: install send2trash to auto-remove embedded figures.")


if __name__ == "__main__":
    main()
