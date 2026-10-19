#!/usr/bin/env python3
"""
setup.py — First-time setup for Research Rolodex.

Run this once before using append-slide.py. It will:
  1. Ask where you want to keep your Research Rolodex folder
  2. Ask for your initial section/theme names
  3. Create the folder structure and subfolders
  4. Create a blank Research Rolodex.pptx with a title slide
  5. Write a starter config.json so you know what fields to fill in

Usage:
    python3 setup.py
"""
import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Inline minimal slide builder (duplicated from append-slide.py to avoid
#    circular imports during first-run setup) ─────────────────────────────────

COLOR_ACCENT   = "221F72"
COLOR_GRAPHITE = "58595B"
COLOR_PAPER    = "F7F7F4"
COLOR_WHITE    = "FFFFFF"
FONT_TITLE     = "Arial Narrow"
FONT_BODY      = "Arial Narrow"

def hex_to_rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _make_title_slide(prs, title, subtitle=""):
    """Create the cover/title slide."""
    layout = prs.slide_layouts[0]
    slide  = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = hex_to_rgb(COLOR_PAPER)

    # Accent bar across the top
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(10), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = hex_to_rgb(COLOR_ACCENT)
    bar.line.fill.background()

    # Main title
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(9.0), Inches(1.2))
    tf = tb.text_frame
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(48)
    r.font.bold = True
    r.font.color.rgb = hex_to_rgb(COLOR_ACCENT)
    r.font.name = FONT_TITLE

    # Subtitle
    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(3.1), Inches(9.0), Inches(0.5))
        tf2 = tb2.text_frame
        p2  = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(18)
        r2.font.color.rgb = hex_to_rgb(COLOR_GRAPHITE)
        r2.font.name = FONT_BODY

    return slide


# ── Setup prompts ─────────────────────────────────────────────────────────────

def prompt(question, default=None):
    if default:
        answer = input(f"{question} [{default}]: ").strip()
        return answer if answer else default
    return input(f"{question}: ").strip()


def main():
    print("=" * 60)
    print("  Research Rolodex — First-Run Setup")
    print("=" * 60)
    print()

    # 1. Rolodex directory
    default_dir = str(Path.home() / "Research Rolodex")
    rolodex_dir = prompt("Where should your Research Rolodex live?", default_dir)
    rolodex_dir = os.path.expanduser(rolodex_dir)

    if os.path.exists(rolodex_dir):
        existing_deck = os.path.join(rolodex_dir, "Research Rolodex.pptx")
        if os.path.exists(existing_deck):
            print(f"\n⚠  A Rolodex already exists at {rolodex_dir}.")
            choice = input("   Overwrite? (y/N): ").strip().lower()
            if choice != 'y':
                print("Setup cancelled. Your existing Rolodex is untouched.")
                sys.exit(0)

    # 2. Deck title
    deck_title = prompt("\nWhat do you want to call your Rolodex?",
                        "Research Rolodex")

    # 3. Section names
    print("\nEnter your initial section/theme names (one per line).")
    print("These are the categories your papers will be sorted into.")
    print("(Press Enter with a blank line when done.)")
    print()
    themes = []
    i = 1
    while True:
        t = input(f"  Theme {i}: ").strip()
        if not t:
            break
        themes.append(t)
        i += 1

    if not themes:
        themes = ["General"]
        print("  (No themes entered — defaulting to 'General')")

    # 4. Optional footer text
    print()
    footer = prompt("Optional footer text (e.g. your lab name, or leave blank):", "")

    # 5. Confirm
    deck_path = os.path.join(rolodex_dir, "Research Rolodex.pptx")
    print()
    print("  Summary:")
    print(f"  Directory:  {rolodex_dir}")
    print(f"  Deck:       {deck_path}")
    print(f"  Sections:   {', '.join(themes)}")
    if footer:
        print(f"  Footer:     {footer}")
    print()
    confirm = input("Create? (Y/n): ").strip().lower()
    if confirm == 'n':
        print("Setup cancelled.")
        sys.exit(0)

    # ── Create directory structure ────────────────────────────────────────────
    os.makedirs(rolodex_dir, exist_ok=True)
    for t in themes:
        safe = t.replace("/", "-").replace("\\", "-").replace(":", "-")
        os.makedirs(os.path.join(rolodex_dir, safe), exist_ok=True)
    print(f"\n✓ Created folder structure in {rolodex_dir}")

    # ── Create the deck ───────────────────────────────────────────────────────
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(5.625)
    _make_title_slide(prs, deck_title)
    prs.save(deck_path)
    print(f"✓ Created {deck_path}")

    # ── Write starter config.json ─────────────────────────────────────────────
    starter_config = {
        "_comment": "Fill in the fields below, then run: python3 append-slide.py config.json",
        "deck_path": deck_path,
        "theme": themes[0],
        "layout": "nofig",
        "footer_text": footer,
        "paper": {
            "title":    "Paper title here",
            "date":     "Jan 2024",
            "authors":  "Smith J, Jones A, et al.",
            "doi":      "10.1234/example",
            "citation": "Journal Name 2024;1(1):1-10",
            "rating":   7,
            "body": [
                {"label": "Problem",    "text": "One-sentence problem statement."},
                {"label": "Methods",    "text": "One-sentence methods summary."},
                {"label": "Results",    "text": "Two sentences with key numbers."},
                {"label": "Conclusion", "text": "One-sentence conclusion."}
            ],
            "notes": (
                "Upload Date: Jan 2024\n\n"
                "My Notes:\n"
                "• \n\n"
                "Highlights:\n"
                "• Key finding or verbatim highlight."
            )
        }
    }
    config_path = os.path.join(rolodex_dir, "example-config.json")
    with open(config_path, 'w') as f:
        json.dump(starter_config, f, indent=2)
    print(f"✓ Wrote example config to {config_path}")

    # ── Write section list (for reference) ───────────────────────────────────
    sections_path = os.path.join(rolodex_dir, "sections.txt")
    with open(sections_path, 'w') as f:
        f.write("# Research Rolodex Sections\n")
        f.write("# One section per line. Order here controls color palette assignment.\n\n")
        for t in themes:
            f.write(t + "\n")
    print(f"✓ Wrote section list to {sections_path}")

    print()
    print("=" * 60)
    print("  Setup complete!")
    print()
    print("  Next steps:")
    print(f"  1. Edit  {config_path}")
    print(f"     (fill in your paper's title, authors, DOI, body text, notes)")
    print(f"  2. Run:  python3 append-slide.py {config_path}")
    print()
    print("  Or drop a PDF into one of the subfolders and let your")
    print("  Claude skill handle the rest (see SKILL.md).")
    print("=" * 60)


if __name__ == "__main__":
    main()
